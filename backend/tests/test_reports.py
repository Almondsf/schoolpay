"""Reports: the dashboard figures, debtors, money received, and downloads."""

from datetime import UTC, date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update

from app.core.pagination import PageParams
from app.core.tenancy import scope_to_school
from app.modules.accounts.models import Role
from app.modules.audit.models import AuditLog
from app.modules.fees.models import Invoice, InvoiceKind
from app.modules.reports import router as reports_router
from app.modules.reports import service
from app.modules.reports.service import AgingBucket, DebtorFilters
from app.modules.students.models import Guardian
from tests.factories import (
    QueryCounter,
    headers_for,
    make_calendar,
    make_levels,
    make_school,
    make_student,
    make_user,
    tenant_writer,
)
from tests.test_fees import add_fees, generate
from tests.test_parents import link_guardian
from tests.test_payments import connect_gateway

NAIRA = 100
DUE = "2026-10-15"
TODAY = date(2026, 10, 20)


async def pay(client, world, student, naira, *, on="2026-10-01", method="cash", reference=""):
    recorded = await client.post(
        "/api/v1/manual-payments",
        headers=world["admin"],
        json={
            "student_id": str(student.id),
            "amount_kobo": naira * NAIRA,
            "method": method,
            "received_on": on,
            "reference": reference,
        },
    )
    assert recorded.status_code == 201, recorded.text
    approved = await client.post(f"/api/v1/manual-payments/{recorded.json()['id']}/approve", headers=world["bursar"])
    assert approved.status_code == 200, approved.text
    return approved.json()


@pytest.fixture
async def world(db, client):
    """
    Three students in three classes (fees 155k, 185k, 205k), bills due 15 October:
    Ada paid in full (and 10k more), Bayo paid 50k, Chidi paid nothing and also
    owes 40k from last session.
    """
    school = make_school(db, "Grace Model College")
    bursar, admin = make_user(db, school, Role.BURSAR), make_user(db, school, Role.SCHOOL_ADMIN)
    _, terms = make_calendar(school)
    levels = make_levels(school)
    add_fees(school, terms[0], levels)
    ada = make_student(school, terms[0], levels[0].arms[0], first="Ada", last="Eze")
    bayo = make_student(school, terms[0], levels[1].arms[0], first="Bayo", last="Ojo")
    chidi = make_student(school, terms[0], levels[2].arms[0], first="Chidi", last="Okafor")
    link_guardian(school, [ada], phone="+2348030000001", name="Mrs Eze")
    bayo_parent = link_guardian(school, [bayo], phone="+2348030000002", name="Mr Ojo")
    link_guardian(school, [chidi], phone="+2348030000003", name='=HYPERLINK("http://evil")')
    connect_gateway(school)
    db.commit()
    w = {
        "school": school,
        "terms": terms,
        "levels": levels,
        "students": {"ada": ada, "bayo": bayo, "chidi": chidi},
        "bayo_parent": bayo_parent,
        "bursar": headers_for(db, bursar),
        "admin": headers_for(db, admin),
    }
    await generate(client, w, due_date=DUE)
    opening = await client.post(
        "/api/v1/invoices",
        headers=w["bursar"],
        json={
            "student_id": str(chidi.id),
            "kind": "opening_balance",
            "lines": [{"description": "Brought forward", "amount_kobo": 40_000 * NAIRA}],
        },
    )
    assert opening.status_code == 201, opening.text
    with tenant_writer() as tw:  # last session's debt was due in July
        tw.execute(
            update(Invoice)
            .where(Invoice.kind == InvoiceKind.OPENING_BALANCE)
            .values(issued_at=datetime(2026, 7, 1, 9, tzinfo=UTC))
        )
    await pay(client, w, ada, 155_000, on="2026-09-20", method="teller", reference="TLR1")
    await pay(client, w, ada, 10_000, on="2026-10-01")
    await pay(client, w, bayo, 50_000, on="2026-10-01")
    return w


def scoped(db, world):
    scope_to_school(db, world["school"].id)
    return db


# --- Overview ---------------------------------------------------------------------------------------


async def test_overview_figures_reconcile(client, world, db):
    response = await client.get("/api/v1/reports/overview", headers=world["bursar"])
    assert response.status_code == 200, response.text
    o = response.json()
    assert o["term_label"] == "2026/2027 First term"
    assert o["invoices"] == 3
    assert o["billed_kobo"] == 545_000 * NAIRA
    assert o["collected_kobo"] == (155_000 + 50_000) * NAIRA
    assert o["outstanding_kobo"] == (135_000 + 205_000) * NAIRA
    assert o["older_outstanding_kobo"] == 40_000 * NAIRA
    assert o["total_outstanding_kobo"] == 380_000 * NAIRA
    assert o["students_owing"] == 2
    assert o["credit_kobo"] == 10_000 * NAIRA  # Ada paid ahead
    assert o["pending_approval"] == 0
    rows = {r["class_name"]: r for r in o["by_class"]}
    assert rows["JSS 1"]["collected_kobo"] == 155_000 * NAIRA and rows["JSS 1"]["students_owing"] == 0
    assert rows["JSS 2"]["outstanding_kobo"] == 135_000 * NAIRA and rows["JSS 2"]["students_owing"] == 1
    assert [r["class_name"] for r in o["by_class"]] == ["JSS 1", "JSS 2", "JSS 3"]  # promotion order


def test_aging_buckets_and_daily_takings(world, db):
    o = service.overview(scoped(db, world), None, today=TODAY)
    aging = {a.bucket: (a.amount_kobo, a.invoices) for a in o.aging}
    assert aging[AgingBucket.DAYS_1_30] == (340_000 * NAIRA, 2)  # this term's bills, 5 days late
    assert aging[AgingBucket.OVER_90] == (40_000 * NAIRA, 1)  # last session's debt, due 1 July
    assert aging[AgingBucket.NOT_DUE] == (0, 0)
    assert [(d.day, d.amount_kobo, d.payments) for d in o.by_day] == [
        (date(2026, 9, 20), 155_000 * NAIRA, 1),
        (date(2026, 10, 1), 60_000 * NAIRA, 2),
    ]
    assert {m.method.value: m.amount_kobo for m in o.by_method} == {
        "teller": 155_000 * NAIRA,
        "cash": 60_000 * NAIRA,
    }
    assert o.received_this_term_kobo == 215_000 * NAIRA


async def test_overview_is_a_fixed_number_of_queries(client, world, db, connection):
    with QueryCounter(connection) as small:
        service.overview(scoped(db, world), None, today=TODAY)
    term, arm = world["terms"][0], world["levels"][0].arms[0]
    for i in range(20):
        make_student(world["school"], term, arm, first=f"S{i}")
    db.commit()
    await generate(client, world, due_date=DUE)
    with QueryCounter(connection) as big:
        service.overview(scoped(db, world), None, today=TODAY)
    assert big.count == small.count


async def test_reversal_counts_against_the_day_it_happened(client, world, db):
    bayo = world["students"]["bayo"]
    payment = await pay(client, world, bayo, 20_000, on="2026-10-05", method="teller", reference="TLR9")
    reversed_ = await client.post(
        f"/api/v1/payments/{payment['payment_id']}/reverse", headers=world["bursar"], json={"reason": "Bounced"}
    )
    assert reversed_.status_code in (200, 201), reversed_.text
    o = service.overview(scoped(db, world), None, today=TODAY)
    assert o.collected_kobo == 205_000 * NAIRA  # back to before the bounced teller
    # The bounced teller doesn't count as Bayo's last payment
    [bayo_row] = [
        d
        for d in service.debtors(db, DebtorFilters(), PageParams(limit=50, offset=0), TODAY).items
        if d.full_name == "Bayo Ojo"
    ]
    assert bayo_row.last_paid_at is not None and bayo_row.last_paid_at.date() == date(2026, 10, 1)


# --- Debtors ----------------------------------------------------------------------------------------


async def debtors(client, world, **query):
    response = await client.get("/api/v1/reports/debtors", headers=world["bursar"], params=query)
    assert response.status_code == 200, response.text
    return response.json()


async def test_debtors_list_filters_and_sorts(client, world):
    everyone = await debtors(client, world)
    assert everyone["total"] == 2 and everyone["total_owed_kobo"] == 380_000 * NAIRA
    chidi, bayo = everyone["items"]
    assert (chidi["full_name"], chidi["owed_kobo"]) == ("Chidi Okafor", 245_000 * NAIRA)
    assert chidi["class_name"] == "JSS 3 A" and chidi["guardian_phone"] == "+2348030000003"
    assert chidi["last_paid_at"] is None
    assert (bayo["full_name"], bayo["owed_kobo"]) == ("Bayo Ojo", 135_000 * NAIRA)

    assert [d["full_name"] for d in (await debtors(client, world, sort="name"))["items"]] == [
        "Bayo Ojo",
        "Chidi Okafor",
    ]
    assert next(d["full_name"] for d in (await debtors(client, world, sort="oldest"))["items"]) == "Chidi Okafor"
    by_class = await debtors(client, world, class_level_id=str(world["levels"][1].id))
    assert [d["full_name"] for d in by_class["items"]] == ["Bayo Ojo"]
    assert by_class["total_owed_kobo"] == 135_000 * NAIRA
    big = await debtors(client, world, min_owed_kobo=200_000 * NAIRA)
    assert [d["full_name"] for d in big["items"]] == ["Chidi Okafor"]
    assert (await debtors(client, world, search="ojo"))["total"] == 1
    assert (await debtors(client, world, search="%"))["total"] == 2  # wildcards are just text
    page = await debtors(client, world, limit=1, offset=1)
    assert page["total"] == 2 and len(page["items"]) == 1


def test_overdue_and_opt_out_shown(world, db):
    with tenant_writer() as tw:
        tw.execute(
            update(Guardian)
            .where(Guardian.id == world["bayo_parent"].id)
            .values(reminders_opt_out_at=datetime.now(UTC))
        )
    session = scoped(db, world)
    before_due = service.debtors(
        session, DebtorFilters(overdue_only=True), PageParams(limit=50, offset=0), date(2026, 10, 1)
    )
    assert [d.full_name for d in before_due.items] == ["Chidi Okafor"]  # only last session's debt was due
    after = service.debtors(session, DebtorFilters(overdue_only=True), PageParams(limit=50, offset=0), TODAY)
    assert after.total == 2
    chidi = next(d for d in after.items if d.full_name == "Chidi Okafor")
    assert chidi.days_overdue == (TODAY - date(2026, 7, 1)).days
    bayo = next(d for d in after.items if d.full_name == "Bayo Ojo")
    assert bayo.reminders_off is True and chidi.reminders_off is False


async def test_debtors_is_a_fixed_number_of_queries(client, world, db, connection):
    with QueryCounter(connection) as small:
        service.debtors(scoped(db, world), DebtorFilters(), PageParams(limit=50, offset=0), TODAY)
    term, arm = world["terms"][0], world["levels"][0].arms[0]
    for i in range(15):
        link_guardian(
            world["school"], [make_student(world["school"], term, arm, first=f"D{i}")], phone=f"+234805000{i:04d}"
        )
    db.commit()
    await generate(client, world, due_date=DUE)
    with QueryCounter(connection) as big:
        page = service.debtors(scoped(db, world), DebtorFilters(), PageParams(limit=50, offset=0), TODAY)
    assert page.total == 17 and big.count == small.count


# --- Money received ---------------------------------------------------------------------------------


async def test_collections_over_a_period(client, world):
    response = await client.get(
        "/api/v1/reports/collections", headers=world["bursar"], params={"start": "2026-09-01", "end": "2026-09-30"}
    )
    c = response.json()
    assert (c["received_kobo"], c["net_kobo"], c["payments"]) == (155_000 * NAIRA, 155_000 * NAIRA, 1)
    whole = await client.get(
        "/api/v1/reports/collections", headers=world["bursar"], params={"start": "2026-09-01", "end": "2026-10-31"}
    )
    assert whole.json()["net_kobo"] == 215_000 * NAIRA
    backwards = await client.get(
        "/api/v1/reports/collections", headers=world["bursar"], params={"start": "2026-10-31", "end": "2026-09-01"}
    )
    assert backwards.status_code == 400
    too_long = await client.get(
        "/api/v1/reports/collections", headers=world["bursar"], params={"start": "2024-01-01", "end": "2026-01-01"}
    )
    assert too_long.status_code == 400


# --- Downloads --------------------------------------------------------------------------------------


async def test_debtors_excel_and_pdf(client, world, db):
    xlsx = await client.get("/api/v1/reports/debtors/export", headers=world["bursar"], params={"format": "xlsx"})
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert xlsx.headers["cache-control"] == "no-store"
    sheet = load_workbook(BytesIO(xlsx.content))["Debtors"]
    assert sheet["A1"].value == "Grace Model College"
    assert [c.value for c in sheet[5]][:4] == ["Student", "Admission no.", "Class", "Owed (NGN)"]
    first = [c.value for c in sheet[6]]
    assert first[0] == "Chidi Okafor" and first[3] == 245_000
    # A parent name that looks like a formula is written as text, never as a formula
    assert first[6] == '\'=HYPERLINK("http://evil")' and sheet.cell(row=6, column=7).data_type == "s"
    assert sheet.cell(row=8, column=4).value == 380_000  # total

    pdf = await client.get("/api/v1/reports/debtors/export", headers=world["bursar"], params={"format": "pdf"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert 'filename="debtors-' in pdf.headers["content-disposition"]

    # Who downloaded parents' contact details is on record (but not the details themselves)
    scope_to_school(db, world["school"].id)
    entries = list(db.scalars(select(AuditLog).where(AuditLog.action == "reports.exported")))
    assert len(entries) == 2 and entries[0].changes["rows"] == 2
    assert "+234" not in str([e.changes for e in entries])


async def test_collections_excel_and_pdf(client, world):
    params = {"start": "2026-09-01", "end": "2026-10-31"}
    xlsx = await client.get(
        "/api/v1/reports/collections/export", headers=world["bursar"], params={**params, "format": "xlsx"}
    )
    book = load_workbook(BytesIO(xlsx.content))
    assert book.sheetnames == ["By day", "By method", "Payments"]
    payments = book["Payments"]
    assert [payments.cell(row=r, column=3).value for r in range(6, 9)] == ["Ada Eze", "Ada Eze", "Bayo Ojo"]
    assert payments.cell(row=6, column=8).value == 155_000
    assert payments.cell(row=6, column=2).value == 1  # receipt number
    pdf = await client.get(
        "/api/v1/reports/collections/export", headers=world["bursar"], params={**params, "format": "pdf"}
    )
    assert pdf.content.startswith(b"%PDF")
    assert 'filename="money-received-2026-09-01-to-2026-10-31.pdf"' in pdf.headers["content-disposition"]


async def test_downloads_are_rate_limited(client, world, monkeypatch):
    monkeypatch.setattr(reports_router, "EXPORT_RATE", "2/hour")
    for _ in range(2):
        assert (await client.get("/api/v1/reports/debtors/export", headers=world["bursar"])).status_code == 200
    assert (await client.get("/api/v1/reports/debtors/export", headers=world["bursar"])).status_code == 429


# --- Access -----------------------------------------------------------------------------------------


async def test_only_staff_who_see_money(client, world):
    for path in ("overview", "debtors", "collections", "debtors/export", "collections/export"):
        assert (await client.get(f"/api/v1/reports/{path}", headers=world["admin"])).status_code == 200  # 2026-10-20
    assert (await client.get("/api/v1/reports/overview")).status_code == 401


async def test_schools_see_only_their_own_money(client, world, db):
    other = make_school(db, "Beta Academy")
    other_bursar = make_user(db, other, Role.BURSAR)
    _, other_terms = make_calendar(other)
    db.commit()
    theirs = headers_for(db, other_bursar)

    their_view = (await client.get("/api/v1/reports/overview", headers=theirs)).json()
    assert their_view["invoices"] == 0 and their_view["credit_kobo"] == 0 and their_view["by_day"] == []
    assert (await client.get("/api/v1/reports/debtors", headers=theirs)).json()["total"] == 0
    peek = await client.get("/api/v1/reports/overview", headers=theirs, params={"term_id": str(world["terms"][0].id)})
    assert peek.status_code == 404
    other_class = await client.get(
        "/api/v1/reports/debtors/export",
        headers=theirs,
        params={"class_level_id": str(world["levels"][0].id)},
    )
    assert other_class.status_code == 404
    money = await client.get(
        "/api/v1/reports/collections", headers=theirs, params={"start": "2026-09-01", "end": "2026-10-31"}
    )
    assert money.json()["payments"] == 0
    # And ours is unaffected
    assert (await client.get("/api/v1/reports/overview", headers=world["bursar"])).json()["invoices"] == 3
    assert other_terms


async def test_no_current_term(client, db):
    school = make_school(db)
    bursar = make_user(db, school, Role.BURSAR)
    make_calendar(school, current_term=0)
    db.commit()
    response = await client.get("/api/v1/reports/overview", headers=headers_for(db, bursar))
    assert response.status_code == 409 and "current term" in response.json()["error"]["message"]
