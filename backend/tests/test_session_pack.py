"""The end-of-session pack: one workbook that closes the school year."""

import uuid
from io import BytesIO

from openpyxl import load_workbook

from app.core.clock import lagos_today
from app.modules.accounts.models import Role
from tests.factories import QueryCounter, headers_for, make_school, make_user
from tests.test_cashbook import take_cash
from tests.test_expenses import spend
from tests.test_payments import generate
from tests.test_payments import world as world  # the fixture

PACK = "/api/v1/accounting/session-pack"
BILLED = 155_000  # naira: JSS 1 tuition + PTA in add_fees


async def download(client, world, **params):
    response = await client.get(PACK, headers=world["proprietor"], params=params)
    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "no-store"
    return response, load_workbook(BytesIO(response.content))


def values(sheet):
    return {row[0]: row[1:] for row in sheet.iter_rows(min_row=6, values_only=True) if row[0]}


async def test_the_pack_closes_the_year(client, world):
    await generate(client, world)
    await take_cash(client, world, 100_000, day=lagos_today())
    expense = await spend(client, world, 30_000, paid_from="bank", name="Rent")
    await client.post(f"/api/v1/expenses/{expense['id']}/approve", headers=world["bursar2"])

    response, book = await download(client, world)
    assert 'filename="end-of-session-2026-2027.xlsx"' in response.headers["content-disposition"]
    assert book.sheetnames == [
        "Summary",
        "By term",
        "Money in",
        "Spending",
        "Debtors",
        "Cash book",
        "Journal",
        "Account totals",
        "Chart of accounts",
    ]
    summary = values(book["Summary"])
    assert summary["Term bills, after discounts"][0] == BILLED
    assert summary["Paid against the term bills"][0] == 100_000
    assert summary["Still owed on the term bills"][0] == BILLED - 100_000
    assert summary["Approved spending in the session"][0] == 30_000
    assert summary["Surplus (deficit if negative)"][0] == 70_000
    terms = values(book["By term"])
    assert terms["First term"][:2] == (1, BILLED) and terms["Second term"][0] == 0
    assert values(book["Spending"])["Rent"][1] == 30_000
    debtors = values(book["Debtors"])
    assert debtors[world["student"].full_name][2] == BILLED - 100_000
    journal = values(book["Account totals"])
    assert journal["Total"][2] == journal["Total"][3]  # debits equal credits


async def test_choosing_a_year_other_schools_and_fixed_queries(client, world, db, connection):
    stranger = headers_for(db, make_user(db, make_school(db, "Beta Academy"), Role.PROPRIETOR))
    db.commit()
    assert (await client.get(PACK, headers=stranger)).status_code == 404  # no school year of its own
    mine = await client.get(PACK, headers=stranger, params={"session_id": str(world["terms"][0].session_id)})
    assert mine.status_code == 404  # not Beta's
    unknown = await client.get(PACK, headers=world["proprietor"], params={"session_id": str(uuid.uuid4())})
    assert unknown.status_code == 404

    await generate(client, world)
    with QueryCounter(connection) as one:
        await download(client, world, session_id=str(world["terms"][0].session_id))
    for naira in (1_000, 2_000, 3_000):
        await take_cash(client, world, naira, day=lagos_today())
        e = await spend(client, world, naira, paid_from="cash")
        await client.post(f"/api/v1/expenses/{e['id']}/approve", headers=world["bursar2"])
    with QueryCounter(connection) as many:
        await download(client, world, session_id=str(world["terms"][0].session_id))
    assert len(many.statements) == len(one.statements)
