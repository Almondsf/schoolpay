"""Fee status by pupil: each fee paid, part paid or unpaid; downloads; isolation; fixed queries."""

from io import BytesIO

from openpyxl import load_workbook

from app.core.tenancy import scope_to_school
from app.modules.accounts.models import Role
from app.modules.reports.fee_status import fee_status
from tests.factories import QueryCounter, headers_for, make_school, make_student, make_user
from tests.test_reports import NAIRA
from tests.test_reports import world as world  # the fixture

URL = "/api/v1/reports/fee-status/export"


def report_for(db, world, **kwargs):
    scope_to_school(db, world["school"].id)
    return fee_status(db, world["terms"][0].id, **kwargs)


def test_each_fee_shows_what_is_left(db, world):
    report = report_for(db, world)
    assert report.scope == "All classes" and report.term_label == "2026/2027 First term"
    assert set(report.columns) == {"PTA", "Tuition"}  # in the order they appear on the bills
    pupils = {p.full_name: p for p in report.pupils}

    ada = pupils["Ada Eze"]
    assert ada.status == "Paid" and all(c.status == "Paid" for c in ada.fees.values())

    bayo = pupils["Bayo Ojo"]  # paid 50,000 of 185,000, not marked for a fee: tuition first (PTA is first on the bill)
    assert bayo.status == "Part paid"
    assert bayo.fees["PTA"].status == "Unpaid"
    assert bayo.fees["Tuition"].status == "Part paid"
    assert bayo.fees["Tuition"].left_kobo == 130_000 * NAIRA
    assert bayo.term_left_kobo == 135_000 * NAIRA and bayo.term_paid_kobo == 50_000 * NAIRA

    chidi = pupils["Chidi Okafor"]  # nothing paid, and owes 40,000 from last session
    assert chidi.status == "Unpaid" and chidi.fees["Tuition"].status == "Unpaid"
    assert chidi.earlier_left_kobo == 40_000 * NAIRA
    assert chidi.total_left_kobo == 245_000 * NAIRA


def test_one_class_only(db, world):
    report = report_for(db, world, class_level_id=world["levels"][1].id)
    assert report.scope == "JSS 2" and [p.full_name for p in report.pupils] == ["Bayo Ojo"]


async def test_excel_has_a_sheet_per_class_and_colours_the_status(client, world):
    response = await client.get(URL, headers=world["bursar"])
    assert response.status_code == 200, response.text
    assert 'filename="fee-status-all-classes-' in response.headers["content-disposition"]
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["All pupils", "JSS 1 A", "JSS 2 A", "JSS 3 A"]
    sheet = book["All pupils"]
    header = [c.value for c in sheet[5]]
    assert header[:3] == ["Student", "Admission no.", "Class"]
    tuition, pta = header.index("Tuition") + 1, header.index("PTA") + 1  # 1-based columns
    assert header[tuition : tuition + 2] == ["Tuition: paid (NGN)", "Tuition: left (NGN)"]
    rows = {sheet.cell(row=r, column=1).value: r for r in range(6, 9)}
    bayo = rows["Bayo Ojo"]  # paid 50,000: all to tuition first
    assert sheet.cell(row=bayo, column=tuition).value == "Part paid"
    assert sheet.cell(row=bayo, column=tuition + 2).value == 130_000
    assert sheet.cell(row=bayo, column=tuition + 1).value + sheet.cell(row=bayo, column=pta + 1).value == 50_000
    assert sheet.cell(row=bayo, column=tuition).fill.fgColor.rgb.endswith("FCEBC7")
    assert "in the order they appear on the bill" in str(sheet.cell(row=sheet.max_row, column=1).value)


async def test_pdf_for_one_class(client, world):
    response = await client.get(
        URL, headers=world["bursar"], params={"format": "pdf", "class_level_id": str(world["levels"][2].id)}
    )
    assert response.status_code == 200 and response.content.startswith(b"%PDF")
    assert 'filename="fee-status-jss-3-' in response.headers["content-disposition"]


async def test_only_staff_who_see_money_and_only_their_school(client, world, db):
    assert (await client.get(URL, headers=world["admin"])).status_code == 200  # school admins see money too
    assert (await client.get(URL)).status_code == 401
    other = make_user(db, make_school(db), Role.BURSAR)
    db.commit()
    stranger = await client.get(URL, headers=headers_for(db, other), params={"term_id": str(world["terms"][0].id)})
    assert stranger.status_code == 404  # another school's term


def test_fixed_number_of_queries(db, world, connection):
    with QueryCounter(connection) as small:
        report_for(db, world)
    for i in range(10):
        make_student(world["school"], world["terms"][0], world["levels"][0].arms[0], first=f"More{i}")
    with QueryCounter(connection) as big:
        report_for(db, world)
    assert big.count == small.count


async def test_a_pupils_breakdown_for_staff(client, world, db):
    bayo = world["students"]["bayo"]
    response = await client.get(f"/api/v1/reports/students/{bayo.id}/fee-breakdown", headers=world["admin"])
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["term_label"] == "2026/2027 First term"
    fees = {f["name"]: f for f in body["fees"]}
    assert fees["PTA"] == {"name": "PTA", "billed_kobo": 5_000 * NAIRA, "paid_kobo": 0, "left_kobo": 5_000 * NAIRA,
                           "status": "Unpaid"}  # fmt: skip
    assert fees["Tuition"]["left_kobo"] == 130_000 * NAIRA and fees["Tuition"]["status"] == "Part paid"
    assert body["total_left_kobo"] == 135_000 * NAIRA and body["earlier_left_kobo"] == 0

    chidi = (await client.get(f"/api/v1/reports/students/{world['students']['chidi'].id}/fee-breakdown",
                              headers=world["bursar"])).json()  # fmt: skip
    assert chidi["earlier_left_kobo"] == 40_000 * NAIRA and chidi["total_left_kobo"] == 245_000 * NAIRA

    other = make_user(db, make_school(db), Role.BURSAR)
    db.commit()
    stranger = await client.get(f"/api/v1/reports/students/{bayo.id}/fee-breakdown", headers=headers_for(db, other))
    assert stranger.status_code == 404


async def test_one_off_charges_show_by_their_own_name(client, world):
    ada = world["students"]["ada"]
    charge = await client.post(
        f"/api/v1/students/{ada.id}/charges",
        headers=world["bursar"],
        json={"description": "School uniform", "unit_kobo": 12_000 * NAIRA},
    )
    assert charge.status_code in (200, 201), charge.text
    body = (await client.get(f"/api/v1/reports/students/{ada.id}/fee-breakdown", headers=world["bursar"])).json()
    names = [f["name"] for f in body["fees"]]
    assert "School uniform" in names and "Other charges" not in names


def _two_classes():
    """JSS 1 pays tuition and a uniform; SSS 3 pays tuition and WAEC. A made-up report, no database."""
    from app.modules.reports.fee_status import FeeCell, FeeStatusReport, PupilStatus

    def pupil(name, klass, fees, earlier=0):
        return PupilStatus(name, name[:3].upper(), klass, {k: FeeCell(*v) for k, v in fees.items()}, {}, earlier)

    pupils = [
        pupil("Ada Eze", "JSS 1 A", {"Tuition": (100_000_00, 0), "Uniform": (15_000_00, 0)}),
        pupil("Bayo Ojo", "JSS 1 A", {"Tuition": (100_000_00, 60_000_00), "Uniform": (15_000_00, 15_000_00)}),
        pupil(
            "Chidi Okafor", "SSS 3 A", {"Tuition": (150_000_00, 150_000_00), "WAEC": (30_000_00, 30_000_00)}, 5_000_00
        ),
        pupil("Dayo Bello", "SSS 3 A", {"Tuition": (150_000_00, 0), "WAEC": (30_000_00, 10_000_00)}),
    ]
    return FeeStatusReport("2026/2027 First term", "All classes", ["Tuition", "Uniform", "WAEC"], pupils)


def _texts(table):
    return [[getattr(cell, "text", cell) for cell in row] for row in table._cellvalues]


def test_each_class_lists_only_its_own_fees_with_what_was_paid():
    from app.modules.reports.exports import _by_class, _fee_status_class_table

    report = _two_classes()
    jss1, sss3 = (_texts(_fee_status_class_table(report.columns, ps)) for ps in _by_class(report.pupils).values())
    assert jss1[0] == ["Pupil", "Tuition", "Uniform", "Paid this term", "Total owed", "Status"]
    assert jss1[2] == [
        "Bayo Ojo",
        "40,000 paid, 60,000 left",
        "Unpaid 15,000",
        "NGN 40,000.00",
        "NGN 75,000.00",
        "Part paid",
    ]
    assert sss3[0] == ["Pupil", "Tuition", "WAEC", "Paid this term", "Earlier debts", "Total owed", "Status"]
    assert sss3[2][1:3] == ["Paid", "20,000 paid, 10,000 left"]
    assert sss3[-1][0] == "Total" and sss3[-1][-2] == "NGN 195,000.00"


def test_pdf_has_a_page_per_class_and_excel_a_sheet_per_class_with_its_fees():
    from reportlab.platypus import PageBreak

    from app.modules.reports import exports

    report = _two_classes()
    assert exports.fee_status_pdf("Demo College", report).startswith(b"%PDF")
    seen = []
    real = exports._pdf
    exports._pdf = lambda school, title, subtitle, blocks: seen.extend(blocks) or real(school, title, subtitle, blocks)
    try:
        exports.fee_status_pdf("Demo College", report)
    finally:
        exports._pdf = real
    assert sum(isinstance(b, PageBreak) for b in seen) == 1  # two classes, two pages

    book = load_workbook(BytesIO(exports.fee_status_xlsx("Demo College", report)))
    assert book.sheetnames == ["All pupils", "JSS 1 A", "SSS 3 A"]
    sss3 = [c.value for c in book["SSS 3 A"][5]]
    assert "WAEC" in sss3 and "Uniform" not in sss3
    assert "Uniform" in [c.value for c in book["All pupils"][5]]
