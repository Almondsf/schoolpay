"""Bulk import of students and guardians from a school's spreadsheet."""

import io
import time

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.core.tenancy import scope_to_school
from app.modules.accounts.models import Role
from app.modules.students.models import Enrollment, Guardian, GuardianStudent, Student
from tests.factories import QueryCounter, headers_for, make_calendar, make_levels, make_school, make_student, make_user

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def school(db):
    a = make_school(db)
    admin = make_user(db, a, Role.SCHOOL_ADMIN)
    _, terms = make_calendar(a)
    levels = make_levels(a, {"JSS 1": ["A", "Gold"], "JSS 2": ["A"], "Primary 4": ["A", "B"]})
    db.commit()
    return {"school": a, "headers": headers_for(db, admin), "terms": terms, "levels": levels}


def xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def upload(client, school, content, *, name="students.xlsx", dry_run=False, mode="skip", consent=True):
    return await client.post(
        "/api/v1/students/import",
        headers=school["headers"],
        files={"file": (name, content, XLSX if name.endswith("xlsx") else "text/csv")},
        data={"dry_run": str(dry_run).lower(), "mode": mode, "consent_confirmed": str(consent).lower()},
    )


def big_sheet(n: int) -> bytes:
    rows: list[list[object]] = [
        ["Adm No", "SURNAME", "Other Names", "Sex", "D.O.B", "Class", "Parent Name", "Parent Phone"]
    ]
    for i in range(n):
        # Siblings: every two students share a parent
        rows.append(
            [
                f"ADM{i:04d}",
                f"Surname{i}",
                f"First{i} Middle{i}",
                "M" if i % 2 else "F",
                "14/03/2014",
                "JSS 1A" if i % 3 else "Primary 4 B",
                f"Parent {i // 2}",
                f"0803{i // 2:07d}",
            ]
        )
    return xlsx(rows)


def counts(db, school_id):
    scope_to_school(db, school_id)
    return {
        model.__name__: db.scalar(select(func.count()).select_from(model))
        for model in (Student, Guardian, GuardianStudent, Enrollment)
    }


async def test_500_students_import_quickly_with_a_fixed_number_of_queries(client, school, db, connection):
    preview = await upload(client, school, big_sheet(500), dry_run=True)
    assert preview.status_code == 200, preview.text
    report = preview.json()
    assert report["errors"] == []
    assert report["students_created"] == 500 and report["guardians_created"] == 250
    assert len(report["preview"]) == 20
    db.expire_all()
    assert counts(db, school["school"].id)["Student"] == 0  # a dry run changes nothing

    started = time.perf_counter()
    with QueryCounter(connection) as big:
        imported = (await upload(client, school, big_sheet(500))).json()
    elapsed = time.perf_counter() - started
    assert imported["students_created"] == 500 and imported["guardian_links_created"] == 500
    assert counts(db, school["school"].id) == {
        "Student": 500,
        "Guardian": 250,
        "GuardianStudent": 500,
        "Enrollment": 500,
    }
    assert elapsed < 30, f"500 rows took {elapsed:.1f}s"

    # Same shape, different size: the number of SQL statements must not grow with the rows
    with QueryCounter(connection) as small:
        await upload(client, school, big_sheet(10).replace(b"ADM", b"ADX"))
    big_selects = [s for s in big.statements if s.lstrip().upper().startswith("SELECT")]
    small_selects = [s for s in small.statements if s.lstrip().upper().startswith("SELECT")]
    assert len(big_selects) == len(small_selects)
    assert big.count <= small.count + 6  # inserts are batched; a few more batches for 500 rows at most


async def test_rows_with_problems_are_reported_and_the_rest_imported(client, school, db):
    content = xlsx(
        [
            ["Admission Number", "Surname", "First Name", "Class", "Arm", "Guardian Phone", "Guardian Name", "Gender"],
            ["A1", "Okafor", "Ada", "JSS 1", "Gold", "08031111111", "Mrs Okafor", "Female"],
            ["A2", "", "Bayo", "JSS 1", "A", "", "", ""],  # no surname
            ["A3", "Bello", "Tunde", "JSS 9", "A", "", "", ""],  # unknown class
            ["A4", "Eze", "Uche", "JSS 1", "A", "12345", "Mr Eze", ""],  # bad phone
            ["A1", "Dup", "Licate", "JSS 1", "A", "", "", ""],  # duplicate admission number
            ["A5", "Musa", "Aisha", "JSS 1", "Blue", "", "", ""],  # unknown arm
            ["A6", "Ibe", "Chika", "jss1", "gold", "", "", "Girl"],  # loose spelling is fine
            ["A7", "Ade", "Sola", "JSS 1", "A", "", "", "Robot"],  # bad gender
        ]
    )
    report = (await upload(client, school, content)).json()
    assert report["students_created"] == 2
    by_row = {e["row"]: e["message"] for e in report["errors"]}
    assert set(by_row) == {3, 4, 5, 6, 7, 9}
    assert "Surname is missing" in by_row[3]
    assert "JSS 9" in by_row[4]
    assert "phone" in by_row[5]
    assert "row 2" in by_row[6]
    assert "Blue" in by_row[7]
    assert "Robot" in by_row[9]


async def test_reimport_skips_existing_or_updates_them(client, school, db):
    first = xlsx([["Admission Number", "Surname", "First Name", "Class"], ["S1", "Okafor", "Ada", "JSS 2"]])
    assert (await upload(client, school, first)).json()["students_created"] == 1
    again = (await upload(client, school, first)).json()
    assert (again["students_created"], again["students_skipped"]) == (0, 1)

    moved = xlsx([["Admission Number", "Surname", "First Name", "Class"], ["S1", "Okafor", "Adaeze", "JSS 1 Gold"]])
    updated = (await upload(client, school, moved, mode="update")).json()
    assert updated["students_updated"] == 1
    student = (await client.get("/api/v1/students?search=S1", headers=school["headers"])).json()["items"][0]
    assert student["first_name"] == "Adaeze"
    assert (student["placement"]["class_level_name"], student["placement"]["arm_name"]) == ("JSS 1", "Gold")


async def test_csv_with_father_and_mother_columns(client, school, db):
    content = (
        b"Reg No,Surname,Firstname,Class,Father Name,Father Phone,Mother Name,Mother Phone\n"
        b"C1,Obi,Ada,JSS 2,Mr Obi,+234 803 222 3333,Mrs Obi,0803-444-5555\n"
    )
    report = (await upload(client, school, content, name="list.csv")).json()
    assert report["students_created"] == 1 and report["guardians_created"] == 2
    student_id = (await client.get("/api/v1/students", headers=school["headers"])).json()["items"][0]["id"]
    detail = (await client.get(f"/api/v1/students/{student_id}", headers=school["headers"])).json()
    relationships = {g["relationship"]: g for g in detail["guardians"]}
    assert relationships["father"]["is_primary"] is True
    assert relationships["mother"]["phone"] == "+2348034445555"


async def test_existing_guardian_is_reused_and_consent_is_required(client, school, db):
    content = xlsx(
        [
            ["Admission Number", "Surname", "First Name", "Class", "Guardian Name", "Guardian Phone"],
            ["G1", "Obi", "Ada", "JSS 2", "Mrs Obi", "08035550000"],
        ]
    )
    refused = await upload(client, school, content, consent=False)
    assert refused.status_code == 400 and refused.json()["error"]["code"] == "consent_required"
    await upload(client, school, content)
    sibling = xlsx(
        [
            ["Admission Number", "Surname", "First Name", "Class", "Guardian Name", "Guardian Phone"],
            ["G2", "Obi", "Emeka", "JSS 2", "Mrs Obi", "2348035550000"],
        ]
    )
    report = (await upload(client, school, sibling)).json()
    assert (report["guardians_created"], report["guardian_links_created"]) == (0, 1)


async def test_bad_files_are_rejected_clearly(client, school):
    missing = xlsx([["Name", "Class"], ["Ada Obi", "JSS 1"]])
    response = await upload(client, school, missing)
    assert response.status_code == 400 and response.json()["error"]["code"] == "missing_columns"
    not_excel = await upload(client, school, b"not really a workbook")
    assert not_excel.status_code == 400
    wrong_type = await upload(client, school, b"x", name="students.pdf")
    assert wrong_type.status_code == 400


async def test_template_lists_the_schools_classes(client, school):
    response = await client.get("/api/v1/students/import-template", headers=school["headers"])
    assert response.status_code == 200 and response.headers["content-type"] == XLSX
    workbook = load_workbook(io.BytesIO(response.content))
    assert next(workbook["Students"].iter_rows(values_only=True))[0] == "Admission Number"
    notes = " ".join(str(r[0]) for r in workbook["How to fill this in"].iter_rows(values_only=True))
    assert "JSS 1 (A, Gold)" in notes


async def test_import_never_touches_another_school(client, school, db):
    other = make_school(db)
    _, other_terms = make_calendar(other)
    other_levels = make_levels(other)
    make_student(other, other_terms[0], other_levels[0].arms[0], admission="SHARED-1", first="Other")
    db.commit()
    content = xlsx([["Admission Number", "Surname", "First Name", "Class"], ["SHARED-1", "Mine", "Ada", "JSS 2"]])
    report = (await upload(client, school, content, mode="update")).json()
    assert report["students_created"] == 1 and report["students_updated"] == 0
    other_term = await client.post(
        "/api/v1/students/import",
        headers=school["headers"],
        files={"file": ("s.xlsx", content, XLSX)},
        data={"term_id": str(other_terms[0].id), "consent_confirmed": "true"},
    )
    assert other_term.status_code == 404
