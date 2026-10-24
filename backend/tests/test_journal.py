"""The accountant's general journal: balanced entries, the school's own account codes, downloads."""

import csv
from io import BytesIO, StringIO

from openpyxl import load_workbook
from sqlalchemy import text

from app.core.clock import lagos_today
from app.core.tenancy import scope_platform_wide
from app.modules.accounts.models import Role
from tests.factories import QueryCounter, headers_for, make_school, make_user
from tests.test_cashbook import bank, take_cash
from tests.test_expenses import spend
from tests.test_payments import NAIRA, generate
from tests.test_payments import world as world  # the fixture

JOURNAL = "/api/v1/accounting/journal"
BILLED = 155_000 * NAIRA  # JSS 1 tuition + PTA in add_fees


async def summary(client, world, headers=None):
    today = lagos_today().isoformat()
    response = await client.get(JOURNAL, headers=headers or world["bursar"], params={"start": today, "end": today})
    assert response.status_code == 200, response.text
    return response.json()


def net(result):
    """{account key: debit - credit} for the period."""
    return {t["account"]["key"]: t["debit_kobo"] - t["credit_kobo"] for t in result["totals"]}


async def test_bills_payments_banking_and_spending_balance(client, world):
    today = lagos_today()
    await generate(client, world)
    await take_cash(client, world, 200_000, day=today)  # ₦155,000 clears the bill, ₦45,000 is paid ahead
    deposit = await bank(client, world, 120_000, day=today)
    await client.post(f"/api/v1/cash-deposits/{deposit['id']}/confirm", headers=world["bursar2"])
    await bank(client, world, 10_000, day=today)  # not yet checked: left out
    expense = await spend(client, world, 30_000, paid_from="bank")
    await client.post(f"/api/v1/expenses/{expense['id']}/approve", headers=world["bursar2"])
    await spend(client, world, 5_000, paid_from="cash")  # not approved: left out

    result = await summary(client, world)
    assert result["debit_kobo"] == result["credit_kobo"] > 0
    assert result["unchecked_deposits"] == 1
    moved = net(result)
    assert moved["cash"] == 80_000 * NAIRA  # 200,000 in, 120,000 banked
    assert moved["bank"] == 90_000 * NAIRA  # 120,000 banked, 30,000 spent
    assert moved["receivable"] == 0  # billed and paid
    assert moved["parent_credit"] == -45_000 * NAIRA
    assert sum(v for k, v in moved.items() if k.startswith("fee:")) == -BILLED
    diesel = next(k for k in moved if k.startswith("expense:"))
    assert moved[diesel] == 30_000 * NAIRA


async def test_voids_reversals_and_credit_applied_on_a_later_day(client, world, db):
    today = lagos_today()
    await take_cash(client, world, 50_000, day=today)  # before any bill: all paid ahead
    # Make the payment look recorded three days ago, so the bill below uses the credit "later"
    scope_platform_wide(db)
    db.execute(text("SET LOCAL session_replication_role = replica"))
    db.execute(text("UPDATE payments SET created_at = created_at - interval '3 days' WHERE amount_kobo > 0"))
    db.execute(text("UPDATE allocations SET created_at = created_at - interval '3 days'"))
    db.execute(text("SET LOCAL session_replication_role = origin"))
    db.commit()
    await generate(client, world)
    result = await summary(client, world)
    moved = net(result)
    assert moved["parent_credit"] == 0  # 50,000 in, then put on the bill
    assert moved["receivable"] == BILLED - 50_000 * NAIRA

    payment_id = db.scalar(text("SELECT id FROM payments WHERE amount_kobo > 0"))
    reversed_ = await client.post(
        f"/api/v1/payments/{payment_id}/reverse", headers=world["bursar"], json={"reason": "Fake notes"}
    )
    assert reversed_.status_code in (200, 201), reversed_.text
    account = (await client.get(f"/api/v1/students/{world['student'].id}/account", headers=world["bursar"])).json()
    voided = await client.post(
        f"/api/v1/invoices/{account['invoices'][0]['id']}/void", headers=world["bursar"], json={"reason": "Wrong class"}
    )
    assert voided.status_code == 200, voided.text
    after = net(await summary(client, world))
    assert after["receivable"] == 0 and after["cash"] == 0  # billed then voided; paid then reversed
    assert all(v == 0 for k, v in after.items() if k.startswith("fee:"))


async def test_the_school_names_its_own_accounts_and_downloads_the_journal(client, world):
    today = lagos_today()
    await generate(client, world)
    await take_cash(client, world, 100_000, day=today, method="teller")
    url = "/api/v1/accounting/accounts"
    taken = await client.put(f"{url}/bank", headers=world["bursar"], json={"code": "1000"})
    assert taken.status_code == 409  # cash has it
    renamed = await client.put(f"{url}/bank", headers=world["bursar"], json={"code": "1200", "name": "GTBank current"})
    assert renamed.json()["code"] == "1200" and renamed.json()["default_code"] == "1010"
    assert (await client.put(f"{url}/nothing", headers=world["bursar"], json={"code": "9"})).status_code == 404

    params = {"start": today.isoformat(), "end": today.isoformat()}
    csv_file = await client.get(f"{JOURNAL}/export", headers=world["proprietor"], params={**params, "format": "csv"})
    assert csv_file.status_code == 200 and csv_file.headers["cache-control"] == "no-store"
    rows = list(csv.DictReader(StringIO(csv_file.content.decode("utf-8-sig"))))
    bank_row = next(r for r in rows if r["Account"] == "GTBank current")
    assert bank_row["Account code"] == "1200" and bank_row["Debit (NGN)"] == "100000.00"
    assert all(world["student"].first_name not in r["Description"] for r in rows)  # admission numbers, no names
    assert sum(float(r["Debit (NGN)"] or 0) for r in rows) == sum(float(r["Credit (NGN)"] or 0) for r in rows)

    xlsx = await client.get(f"{JOURNAL}/export", headers=world["proprietor"], params=params)
    book = load_workbook(BytesIO(xlsx.content))
    assert book.sheetnames == ["Journal", "Account totals", "Chart of accounts"]

    back = await client.put(f"{url}/bank", headers=world["bursar"], json={"code": "", "name": ""})
    assert back.json()["code"] == "1010" and back.json()["name"] == "Bank"


async def test_other_schools_and_a_fixed_number_of_queries(client, world, db, connection):
    stranger = headers_for(db, make_user(db, make_school(db, "Beta Academy"), Role.BURSAR))
    db.commit()
    await client.put("/api/v1/accounting/accounts/bank", headers=stranger, json={"name": "Beta bank"})
    await generate(client, world)
    await take_cash(client, world, 10_000, day=lagos_today())
    assert (await summary(client, world, stranger))["entries"] == 0
    accounts = (await client.get("/api/v1/accounting/accounts", headers=world["bursar"])).json()
    assert next(a for a in accounts if a["key"] == "bank")["name"] == "Bank"  # Beta's name is Beta's
    with QueryCounter(connection) as one:
        await summary(client, world)
    for _ in range(3):
        await take_cash(client, world, 1_000, day=lagos_today())
        e = await spend(client, world, 1_000, paid_from="bank")
        await client.post(f"/api/v1/expenses/{e['id']}/approve", headers=world["bursar2"])
    with QueryCounter(connection) as many:
        await summary(client, world)
    assert len(many.statements) == len(one.statements)
    too_long = await client.get(JOURNAL, headers=world["bursar"], params={"start": "2025-01-01", "end": "2026-10-24"})
    assert too_long.status_code == 400
