"""The journal as an Excel workbook (journal, account totals, chart) or a flat CSV for accounting software."""

import csv
from io import StringIO

from openpyxl import Workbook

from app.modules.reports.exports import _day, _save, _sheet, naira, safe_text

from .journal import Journal, lines

CSV = "text/csv"
JOURNAL_HEADERS = [
    ("Date", 13),
    ("Entry", 9),
    ("Reference", 16),
    ("Account code", 13),
    ("Account", 32),
    ("Description", 44),
    ("Debit (NGN)", 15),
    ("Credit (NGN)", 15),
]


def journal_xlsx(school: str, j: Journal) -> bytes:
    period = f"{_day(j.start)} to {_day(j.end)}"
    wb = Workbook()
    sheet = wb.create_sheet("Journal", 0)
    wb.remove(wb.worksheets[1])  # the default empty sheet
    _sheet(
        sheet,
        school=school,
        title="General journal",
        subtitle=period,
        headers=JOURNAL_HEADERS,
        rows=[
            [
                e.day,
                e.number,
                e.reference,
                a.code,
                a.name,
                e.description,
                naira(dr) if dr else None,
                naira(cr) if cr else None,
            ]
            for e, a, dr, cr in lines(j)
        ],
        money_columns={7, 8},
        total=["Total", "", "", "", "", "", naira(j.debit_kobo), naira(j.credit_kobo)],
    )
    for (cell,) in sheet.iter_rows(min_row=6, min_col=1, max_col=1):
        if cell.value is not None and not isinstance(cell.value, str):
            cell.number_format = "dd mmm yyyy"
    _sheet(
        wb.create_sheet("Account totals"),
        school=school,
        title="Account totals for the period",
        subtitle=period,
        headers=[("Account code", 13), ("Account", 32), ("Type", 11), ("Debit (NGN)", 15), ("Credit (NGN)", 15)],
        rows=[
            [t.account.code, t.account.name, t.account.type, naira(t.debit_kobo), naira(t.credit_kobo)]
            for t in j.totals
        ],
        money_columns={4, 5},
        total=["Total", "", "", naira(j.debit_kobo), naira(j.credit_kobo)],
    )
    _sheet(
        wb.create_sheet("Chart of accounts"),
        school=school,
        title="Chart of accounts",
        subtitle="WardPay's default codes, with the school's own where it has set them",
        headers=[("Account code", 13), ("Account", 32), ("Type", 11)],
        rows=[[a.code, a.name, a.type] for a in j.accounts],
        money_columns=set(),
    )
    return _save(wb)


def journal_csv(j: Journal) -> bytes:
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow([h for h, _ in JOURNAL_HEADERS])
    for e, a, dr, cr in lines(j):
        writer.writerow(
            [
                e.day.isoformat(),
                e.number,
                safe_text(e.reference),
                safe_text(a.code),
                safe_text(a.name),
                safe_text(e.description),
                f"{naira(dr):.2f}" if dr else "",
                f"{naira(cr):.2f}" if cr else "",
            ]
        )
    return out.getvalue().encode("utf-8-sig")  # the byte-order mark makes Excel read it as UTF-8
